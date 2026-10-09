from io import BytesIO

from openpyxl import Workbook

from imports.services.schema import CORE_SHEETS


def build_test_workbook(include_all_core=True, invalid_source=False):
    wb = Workbook()
    wb.remove(wb.active)

    validation = wb.create_sheet("Validation Lists")
    validation.append(["VALIDATION LISTS"])
    validation.append([])
    validation.append([])
    validation.append([
        "Data Type", "Data Structure", "Source Level", "Publish Cadence", "File Format",
        "Asset Status", "Evidence Level", "Rubric Layer", "Scope", "Data Confidence",
        "Policy Category", "Current Status", "Impact Confidence", "Availability Status",
        "Comparability Status", "Evidence Confidence", "Definition Change Flag", "EIM Coverage",
    ])
    validation.append([
        "Government Dataset", "Structured", "State", "Annual", "XLSX", "Live", "Level 1",
        "1. CPS Performance", "Minnesota", "Verified Fact", "Funding", "Active", "Observed",
        "Available", "Comparable", "High", "No", "Fully Covered",
    ])

    sheets = list(CORE_SHEETS)
    if not include_all_core:
        sheets.remove("Metric Comparability")

    rows = {
        "Master Data Sources": ["DS-001", "Minnesota source", "MN agency", "Government Dataset", "Minnesota", "2024", "High", "https://example.org/source", "BF-001", "", "Tester", "Structured", "", "State", "Minnesota", "Annual", "Official page", "", "2026-10-01"],
        "Source Assets": ["RA-001", "DS-999" if invalid_source else "DS-001", "Annual report", "https://example.org/report.xlsx", "XLSX", 2024, "2024", "Minnesota", "Live", "2026-10-01", "Foster care population", "CPS", "BF-001", "Level 1", "1. CPS Performance", "Tester", "2026-10-01", "Verified", ""],
        "Master Observation Data": ["Minnesota", "BF-001", "Foster care population", 2024, 100, "https://example.org/report.xlsx", "Verified Fact", "", "RA-001"],
        "Metric Comparability": ["BF-001", "Foster care population", "Minnesota", "RA-001", "2024", "Children in foster care", "Count", "Child population", "Children", "Annual", "State", "No", "Complete", "Available", "Comparable", "", "High", "Tester", "2026-10-01", "1.2", ""],
        "Rubric Variable Map": ["RV-001", "1. CPS Performance", "Reporting", "Reports and report rate", "System activity", "Fully Covered", "BF-001", "DS-001", "", "Tester", "2026-10-01", ""],
    }

    headers = {
        "Master Data Sources": ["Source ID", "Source Name", "Organization", "Data Type", "Geographic Scope", "Time Range", "Reliability", "URL", "Fields Covered", "Access Notes", "Added By", "Data Structure", "Parent Source ID", "Source Level", "State/Jurisdiction", "Publish Cadence", "Cadence Evidence", "Publishing Page ID", "Last Checked Date"],
        "Source Assets": ["Report Asset ID", "Source ID", "Report Title", "Report Asset URL", "File Format", "Publication Year", "Years Covered", "Relevant Geography", "Asset Status", "Last Checked Date", "Published Columns or Measures", "Subject", "Used By (EIM Rows)", "Evidence Level", "Rubric Layer", "Reviewer", "Review Date", "Verification Note", "Notes / Open Questions"],
        "Master Observation Data": ["Scope", "Field ID", "Field Name", "Year", "Value", "Source URL", "Data Confidence", "Notes", "Report Asset ID"],
        "Metric Comparability": ["Field ID", "Field Name", "Scope", "Source Asset IDs Used", "Years With Values", "Metric Definition", "Numerator Definition", "Denominator Definition", "Population Scope", "Reporting Period", "Geography", "Definition Change Flag", "Data Completeness", "Availability Status", "Comparability Status", "Comparability Note", "Evidence Confidence", "Reviewer", "Review Date", "Framework Version", "Unresolved Question"],
        "Rubric Variable Map": ["Variable ID", "Rubric Layer", "Variable Group", "Candidate Variable", "Working Definition", "EIM Coverage", "Mapped EIM Fields", "Candidate Source IDs", "MFL Mapping", "Reviewer", "Review Date", "Notes"],
    }

    for name in sheets:
        ws = wb.create_sheet(name)
        ws.append([name])
        ws.append(["Test description"])
        ws.append([])
        ws.append(headers[name])
        ws.append(rows[name])

    schema = wb.create_sheet("Observation Data Schema")
    schema.append(["Observation Data Schema"])
    schema.append([])
    schema.append([])
    schema.append(["Metric_ID", "Baseline Field", "Category", "Metric Name", "Unit", "Definition", "Lane", "Priority", "Prepared Dataset Folder"])
    schema.append(["COST_001", "BF-001", "Direct System", "Foster care population", "count", "Children in foster care", "Lane 3", "Baseline", ""])

    output = BytesIO()
    wb.save(output)
    return output.getvalue()
