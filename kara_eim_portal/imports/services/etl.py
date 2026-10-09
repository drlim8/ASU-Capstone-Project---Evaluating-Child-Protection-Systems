import re
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from django.db import transaction
from openpyxl import load_workbook

from imports.models import AuditEvent, DatasetImport, StagedRecord

from .schema import CLUSTER_COLUMNS, CORE_SHEETS, ID_PATTERNS, PICKLIST_COLUMNS, SUPPORTED_SHEETS


def is_blank(value):
    return value is None or (isinstance(value, str) and not value.strip())


def normalize_value(value):
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, float) and value.is_integer():
        return int(value)
    if isinstance(value, str):
        return value.strip()
    return value


def issue(severity, code, message, sheet="", row=None, column=""):
    return {
        "severity": severity,
        "code": code,
        "sheet": sheet,
        "row": row,
        "column": column,
        "message": message,
    }


def _header_values(ws, header_row):
    raw = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))
    headers = []
    for value in raw:
        if value is None:
            break
        headers.append(str(value).strip())
    return headers


def _read_picklists(workbook):
    if "Validation Lists" not in workbook.sheetnames:
        return {}
    ws = workbook["Validation Lists"]
    headers = _header_values(ws, 4)
    values = {header: set() for header in headers}
    for row in ws.iter_rows(min_row=5, values_only=True):
        for index, header in enumerate(headers):
            if index < len(row) and not is_blank(row[index]):
                values[header].add(str(normalize_value(row[index])))
    return values


def _record_key(record, columns):
    parts = [str(record.get(column, "")).strip() for column in columns]
    return " | ".join(parts)


def _split_ids(value):
    if is_blank(value):
        return []
    return [part.strip() for part in re.split(r"[;,]", str(value)) if part.strip()]


def inspect_workbook(source):
    workbook_source = Path(source) if isinstance(source, (str, Path)) else source
    if hasattr(workbook_source, "seek"):
        workbook_source.seek(0)
    if not zipfile.is_zipfile(workbook_source):
        raise ValueError("The uploaded file is not a valid .xlsx workbook.")
    if hasattr(workbook_source, "seek"):
        workbook_source.seek(0)

    workbook = load_workbook(workbook_source, read_only=True, data_only=True)
    picklists = _read_picklists(workbook)
    errors = []
    warnings = []
    records = []
    stats = {}
    clusters = defaultdict(Counter)

    for sheet_name in CORE_SHEETS:
        if sheet_name not in workbook.sheetnames:
            errors.append(issue("error", "missing_sheet", f"Required sheet '{sheet_name}' is missing.", sheet=sheet_name))

    for sheet_name, config in SUPPORTED_SHEETS.items():
        if sheet_name not in workbook.sheetnames:
            continue
        ws = workbook[sheet_name]
        headers = _header_values(ws, config["header_row"])
        missing_columns = [column for column in config["required_columns"] if column not in headers]
        for column in missing_columns:
            errors.append(
                issue("error", "missing_column", f"Required column '{column}' is missing.", sheet=sheet_name, column=column)
            )
        if missing_columns:
            continue

        seen_keys = {}
        sheet_total = 0
        sheet_valid = 0
        sheet_invalid = 0
        for row_number, row in enumerate(
            ws.iter_rows(min_row=config["header_row"] + 1, max_col=len(headers), values_only=True),
            start=config["header_row"] + 1,
        ):
            if all(is_blank(value) for value in row):
                continue
            payload = {headers[index]: normalize_value(row[index]) for index in range(len(headers))}
            nonblank_count = sum(not is_blank(value) for value in payload.values())
            if nonblank_count == 1:
                # The source workbook uses one-cell section labels, notes, and preallocated placeholder IDs.
                continue
            row_issues = []
            sheet_total += 1

            for column in config["required_values"]:
                if is_blank(payload.get(column)):
                    row_issues.append(
                        issue("error", "required_value", f"'{column}' is required.", sheet_name, row_number, column)
                    )

            for column in set(config["required_columns"]) - set(config["required_values"]):
                if is_blank(payload.get(column)):
                    row_issues.append(
                        issue(
                            "warning",
                            "recommended_value",
                            f"'{column}' is blank and should be reviewed.",
                            sheet_name,
                            row_number,
                            column,
                        )
                    )

            key = _record_key(payload, config["key_columns"])
            if key and not all(is_blank(payload.get(column)) for column in config["key_columns"]):
                if key in seen_keys:
                    row_issues.append(
                        issue(
                            "error",
                            "duplicate_key",
                            f"Duplicate key '{key}'. First seen on row {seen_keys[key]}.",
                            sheet_name,
                            row_number,
                        )
                    )
                else:
                    seen_keys[key] = row_number

            for column, pattern in ID_PATTERNS.items():
                value = payload.get(column)
                if not is_blank(value) and not re.match(pattern, str(value)):
                    row_issues.append(
                        issue("error", "invalid_identifier", f"'{value}' does not match the expected format.", sheet_name, row_number, column)
                    )

            for column, list_name in PICKLIST_COLUMNS.items():
                value = payload.get(column)
                allowed = picklists.get(list_name, set())
                if not is_blank(value) and allowed and str(value) not in allowed:
                    row_issues.append(
                        issue(
                            "error",
                            "invalid_picklist_value",
                            f"'{value}' is not an approved {list_name} value.",
                            sheet_name,
                            row_number,
                            column,
                        )
                    )

            if sheet_name == "Master Observation Data" and is_blank(payload.get("Source URL")):
                row_issues.append(
                    issue("warning", "missing_source", "Observation has no Source URL.", sheet_name, row_number, "Source URL")
                )

            for column in CLUSTER_COLUMNS:
                value = payload.get(column)
                if not is_blank(value):
                    clusters[column][str(value)] += 1

            is_valid = not any(item["severity"] == "error" for item in row_issues)
            if is_valid:
                sheet_valid += 1
            else:
                sheet_invalid += 1
            for item in row_issues:
                (errors if item["severity"] == "error" else warnings).append(item)
            records.append(
                {
                    "sheet_name": sheet_name,
                    "source_row": row_number,
                    "record_key": key,
                    "payload": payload,
                    "is_valid": is_valid,
                    "issues": row_issues,
                }
            )

        stats[sheet_name] = {"total": sheet_total, "valid": sheet_valid, "invalid": sheet_invalid}

    _validate_relationships(records, errors, warnings)
    error_locations = {(e["sheet"], e["row"]) for e in errors if e.get("row")}
    for record in records:
        if (record["sheet_name"], record["source_row"]) in error_locations:
            record["is_valid"] = False
            related = [
                item
                for item in errors
                if item.get("sheet") == record["sheet_name"] and item.get("row") == record["source_row"]
            ]
            existing = {(item["code"], item.get("column"), item["message"]) for item in record["issues"]}
            record["issues"].extend(
                item for item in related if (item["code"], item.get("column"), item["message"]) not in existing
            )

    for sheet_name, sheet_stats in stats.items():
        sheet_records = [record for record in records if record["sheet_name"] == sheet_name]
        sheet_stats["valid"] = sum(record["is_valid"] for record in sheet_records)
        sheet_stats["invalid"] = sum(not record["is_valid"] for record in sheet_records)

    return {
        "sheets_found": workbook.sheetnames,
        "processed_sheets": list(stats),
        "stats": stats,
        "records": records,
        "errors": errors,
        "warnings": warnings,
        "clusters": {column: dict(counter.most_common()) for column, counter in clusters.items()},
    }


def _validate_relationships(records, errors, warnings):
    by_sheet = defaultdict(list)
    for record in records:
        by_sheet[record["sheet_name"]].append(record)

    source_ids = {r["payload"].get("Source ID") for r in by_sheet["Master Data Sources"]}
    asset_ids = {r["payload"].get("Report Asset ID") for r in by_sheet["Source Assets"]}
    field_ids = {r["payload"].get("Baseline Field") for r in by_sheet["Observation Data Schema"]}

    for record in by_sheet["Source Assets"]:
        value = record["payload"].get("Source ID")
        if not is_blank(value) and value not in source_ids:
            errors.append(
                issue("error", "unknown_source", f"Source ID '{value}' is not in Master Data Sources.", record["sheet_name"], record["source_row"], "Source ID")
            )

    for record in by_sheet["Master Observation Data"]:
        payload = record["payload"]
        field_id = payload.get("Field ID")
        asset_id = payload.get("Report Asset ID")
        if field_ids and not is_blank(field_id) and field_id not in field_ids:
            errors.append(
                issue("error", "unknown_field", f"Field ID '{field_id}' is not in Observation Data Schema.", record["sheet_name"], record["source_row"], "Field ID")
            )
        if not is_blank(asset_id) and asset_id not in asset_ids:
            errors.append(
                issue("error", "unknown_asset", f"Report Asset ID '{asset_id}' is not in Source Assets.", record["sheet_name"], record["source_row"], "Report Asset ID")
            )

    for record in by_sheet["Metric Comparability"]:
        payload = record["payload"]
        field_id = payload.get("Field ID")
        if field_ids and not is_blank(field_id) and field_id not in field_ids:
            errors.append(
                issue("error", "unknown_field", f"Field ID '{field_id}' is not in Observation Data Schema.", record["sheet_name"], record["source_row"], "Field ID")
            )
        for asset_id in _split_ids(payload.get("Source Asset IDs Used")):
            if asset_id not in asset_ids:
                warnings.append(
                    issue("warning", "unknown_asset", f"Source asset '{asset_id}' is not in Source Assets.", record["sheet_name"], record["source_row"], "Source Asset IDs Used")
                )


def process_dataset_import(dataset_import, actor=None):
    dataset_import.status = DatasetImport.Status.VALIDATING
    dataset_import.save(update_fields=["status"])
    dataset_import.staged_records.all().delete()

    try:
        dataset_import.uploaded_file.open("rb")
        try:
            result = inspect_workbook(dataset_import.uploaded_file.file)
        finally:
            dataset_import.uploaded_file.close()
        with transaction.atomic():
            staged = [StagedRecord(dataset_import=dataset_import, **record) for record in result["records"]]
            StagedRecord.objects.bulk_create(staged, batch_size=500)
            total_records = len(staged)
            invalid_records = sum(not record.is_valid for record in staged)
            dataset_import.selected_sheets = result["processed_sheets"]
            dataset_import.summary = {
                "total_records": total_records,
                "valid_records": total_records - invalid_records,
                "invalid_records": invalid_records,
                "error_count": len(result["errors"]),
                "warning_count": len(result["warnings"]),
                "sheet_stats": result["stats"],
            }
            dataset_import.validation_errors = result["errors"]
            dataset_import.validation_warnings = result["warnings"]
            dataset_import.cluster_summary = result["clusters"]
            dataset_import.status = DatasetImport.Status.NEEDS_CORRECTION if result["errors"] else DatasetImport.Status.READY
            dataset_import.save(
                update_fields=[
                    "selected_sheets",
                    "summary",
                    "validation_errors",
                    "validation_warnings",
                    "cluster_summary",
                    "status",
                ]
            )
            AuditEvent.objects.create(
                dataset_import=dataset_import,
                actor=actor,
                action="validated",
                details={"errors": len(result["errors"]), "warnings": len(result["warnings"]), "records": total_records},
            )
        return result
    except Exception as exc:
        dataset_import.status = DatasetImport.Status.FAILED
        dataset_import.validation_errors = [issue("error", "processing_failed", str(exc))]
        dataset_import.save(update_fields=["status", "validation_errors"])
        AuditEvent.objects.create(
            dataset_import=dataset_import,
            actor=actor,
            action="processing_failed",
            details={"message": str(exc)},
        )
        raise
