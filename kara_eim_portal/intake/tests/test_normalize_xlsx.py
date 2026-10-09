import tempfile
import unittest
from pathlib import Path

from imports.tests.workbook_factory import build_test_workbook
from intake.services.normalizers.xlsx import normalize_xlsx
from intake.services.types import NormalizeError
from intake.tests import factories


class NormalizeXlsxTests(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)

    def _write(self, data: bytes) -> Path:
        p = Path(self._tmp.name) / "book.xlsx"
        p.write_bytes(data)
        return p

    def _links(self):
        return normalize_xlsx(self._write(factories.link_workbook()))

    def test_each_sheet_is_a_page_with_table(self):
        result = self._links()
        self.assertEqual(result.pages[0].label, "Source Assets")
        self.assertEqual(
            result.pages[0].tables[0].rows[0],
            ["Report Asset ID", "Report Asset URL", "Notes"],
        )

    def test_blank_cells_are_none_not_zero(self):
        rows = self._links().pages[0].tables[0].rows
        self.assertIsNone(rows[2][2])
        self.assertIsNone(rows[3][1])

    def test_formula_cell_uses_cached_value_or_none(self):
        result = self._links()
        rows = result.pages[0].tables[0].rows
        self.assertIsNone(rows[3][2])
        self.assertNotIn("=SUM", result.pages[0].text)
        for row in rows:
            for c in row:
                self.assertFalse(c and "=SUM" in c)

    def test_url_column_produces_candidate_links(self):
        links = self._links().links
        self.assertEqual(len(links), 2)
        self.assertEqual(links[0].label, "Source Assets!B2")
        self.assertEqual(links[0].context, "Report Asset URL")

    def test_url_column_detected_by_content_without_header(self):
        result = normalize_xlsx(self._write(factories.where_workbook()))
        self.assertEqual(len(result.links), 3)
        self.assertEqual(result.links[0].context, "Where")

    def test_eim_workbook_warning(self):
        result = normalize_xlsx(self._write(build_test_workbook()))
        self.assertIn("eim_workbook", [w.code for w in result.warnings])

    def test_damaged_workbook_raises(self):
        with self.assertRaises(NormalizeError):
            normalize_xlsx(self._write(factories.link_workbook()[:100]))

    def test_embedded_image_is_extracted(self):
        import io

        import openpyxl
        from openpyxl.drawing.image import Image as XlImage

        wb = openpyxl.Workbook()
        wb.active["A1"] = "x"
        wb.active.add_image(XlImage(io.BytesIO(factories.png_bytes(100, 80))), "B2")
        buf = io.BytesIO()
        wb.save(buf)
        result = normalize_xlsx(self._write(buf.getvalue()))
        img = result.pages[0].images[0]
        self.assertEqual((img.ext, img.width, img.height), ("png", 100, 80))
