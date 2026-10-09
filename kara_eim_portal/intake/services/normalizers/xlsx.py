"""XLSX normalizer: one page/table per worksheet, images, URL-column links. No Django imports."""
from __future__ import annotations

import io
import zipfile
from pathlib import Path

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException
from PIL import Image

from imports.services.schema import CORE_SHEETS, SUPPORTED_SHEETS
from intake.services.types import (
    ImageData,
    IntakeWarning,
    LinkData,
    NormalizedResult,
    NormalizeError,
    PageData,
    TableData,
)

DAMAGED_MESSAGE = "This workbook is damaged and could not be read."
EIM_MESSAGE = "This looks like an EIM workbook. You can also load it through the EIM import."
URL_HEADER_WORDS = ("url", "link")
URL_CONTENT_THRESHOLD = 0.5
HEADER_SCAN_ROWS = 10


def _text(value):
    if value is None:
        return None
    s = str(value)
    return s if s != "" else None


def _is_url(value) -> bool:
    return isinstance(value, str) and value.strip().lower().startswith(("http://", "https://"))


def _trim(rows: list[list]) -> list[list]:
    """Trim trailing blank rows and columns (the used range)."""
    while rows and all(c is None for c in rows[-1]):
        rows.pop()
    width = max((i + 1 for r in rows for i, c in enumerate(r) if c is not None), default=0)
    return [r[:width] for r in rows]


def _extract_images(ws, page_no: int, warnings: list[IntakeWarning]) -> list[ImageData]:
    images: list[ImageData] = []
    for img in getattr(ws, "_images", []):
        try:
            data = img._data()
            with Image.open(io.BytesIO(data)) as pil:
                fmt = (pil.format or "PNG").lower()
                width, height = pil.size
            images.append(
                ImageData(
                    index=len(images),
                    content=data,
                    ext="jpg" if fmt == "jpeg" else fmt,
                    width=width,
                    height=height,
                )
            )
        except Exception:  # noqa: BLE001 - one bad image must not fail the workbook
            warnings.append(
                IntakeWarning(
                    "image_unreadable",
                    f"An image on sheet {page_no} could not be read and was skipped.",
                    page=page_no,
                )
            )
    return images


def _non_blank(row: list) -> int:
    return sum(c is not None for c in row)


def _header_index(sheet_name: str, rows: list[list]) -> int | None:
    """Index in `rows` of the header row.

    KARA sheets open with a title, a description and a blank row, so the header
    is not simply the first non-empty row. Known EIM sheets use the schema's
    header row when that row holds at least one expected column name. Otherwise
    the header is the first of the first HEADER_SCAN_ROWS rows whose non-blank
    count is at least max(2, half the widest of those rows); failing that, the
    first non-empty row.
    """
    spec = SUPPORTED_SHEETS.get(sheet_name)
    if spec and spec.get("header_row"):
        idx = spec["header_row"] - 1
        if idx < len(rows) and set(spec.get("required_columns", ())) & {
            c.strip() for c in rows[idx] if c is not None
        }:
            return idx
    window = rows[:HEADER_SCAN_ROWS]
    widest = max((_non_blank(r) for r in window), default=0)
    threshold = max(2, widest / 2)
    for i, r in enumerate(window):
        if _non_blank(r) >= threshold:
            return i
    return next((i for i, r in enumerate(rows) if _non_blank(r)), None)


def _url_columns(rows: list[list], header_idx: int | None) -> set[int]:
    """Set of 0-based URL column indexes, judged by header word or by content."""
    if header_idx is None:
        return set()
    header = rows[header_idx]
    cols: set[int] = set()
    for c in range(len(header)):
        h = header[c]
        if h is not None and any(w in h.lower() for w in URL_HEADER_WORDS):
            cols.add(c)
            continue
        cells = [r[c] for r in rows[header_idx + 1:] if c < len(r) and r[c] is not None]
        if cells and sum(_is_url(v) for v in cells) / len(cells) >= URL_CONTENT_THRESHOLD:
            cols.add(c)
    return cols


def _links(ws, rows: list[list], header_idx: int | None, seen: set[str]) -> list[LinkData]:
    links: list[LinkData] = []
    if header_idx is None:
        return links
    cols = _url_columns(rows, header_idx)
    for c in sorted(cols):
        context = rows[header_idx][c] or ""
        for r in range(header_idx + 1, len(rows)):
            value = rows[r][c] if c < len(rows[r]) else None
            if not _is_url(value):
                continue
            url = value.strip()
            if url in seen:
                continue
            seen.add(url)
            coord = ws.cell(row=r + 1, column=c + 1).coordinate
            links.append(LinkData(url=url, label=f"{ws.title}!{coord}", context=context))
    return links


def normalize_xlsx(path: Path) -> NormalizedResult:
    try:
        wb = openpyxl.load_workbook(str(path), data_only=True)
    except (zipfile.BadZipFile, InvalidFileException) as exc:
        raise NormalizeError(DAMAGED_MESSAGE) from exc
    except Exception as exc:  # noqa: BLE001 - damaged archives surface many error types
        raise NormalizeError(DAMAGED_MESSAGE) from exc
    warnings: list[IntakeWarning] = []
    pages: list[PageData] = []
    links: list[LinkData] = []
    seen: set[str] = set()
    try:
        for n, ws in enumerate(wb.worksheets, start=1):
            rows = _trim([[_text(c) for c in row] for row in ws.iter_rows(values_only=True)])
            header_idx = _header_index(ws.title, rows)
            header = list(rows[header_idx]) if header_idx is not None else None
            tables = [TableData(index=0, rows=rows, header_guess=header)]
            text = "\n".join("\t".join(c for c in r if c is not None) for r in rows)
            text = "\n".join(line for line in text.split("\n") if line)
            pages.append(
                PageData(
                    number=n,
                    text=text,
                    has_text_layer=True,
                    label=ws.title,
                    tables=tables,
                    images=_extract_images(ws, n, warnings),
                )
            )
            links.extend(_links(ws, rows, header_idx, seen))
        if set(CORE_SHEETS) <= set(wb.sheetnames):
            warnings.append(IntakeWarning("eim_workbook", EIM_MESSAGE))
    finally:
        wb.close()
    return NormalizedResult(pages=pages, links=links, warnings=warnings)
